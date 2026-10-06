"""Public Expo registration and administrator-only attendance reports."""
import io
from xml.sax.saxutils import escape

import qrcode
from email_validator import validate_email, EmailNotValidError
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, current_app, abort
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from app import db
from app.models import ExpoAttendance
from app.routes import roles_required
from app.time_utils import local_datetime

expo = Blueprint('expo', __name__)
CATEGORIES = {'estudiante': 'Estudiante', 'catedratico': 'Catedrático', 'publico': 'Público general'}


@expo.route('/expo-san-pablo', methods=['GET', 'POST'])
def registration():
    values = request.form if request.method == 'POST' else {}
    if request.method == 'POST':
        try:
            name = request.form.get('name', '').strip()
            if not 2 <= len(name) <= 120:
                raise ValueError('Escribe tu nombre completo (entre 2 y 120 caracteres).')
            try:
                email = validate_email(request.form.get('email', '').strip(), check_deliverability=False).normalized.casefold()
            except EmailNotValidError:
                raise ValueError('Escribe un correo electrónico válido, personal o institucional.')
            if len(email) > 254:
                raise ValueError('El correo es demasiado largo.')
            category = request.form.get('category', '')
            if category not in CATEGORIES:
                raise ValueError('Selecciona estudiante, catedrático o público general.')
            participation = request.form.get('participation', '').strip()
            opinion = request.form.get('opinion', '').strip()
            if not 1 <= len(participation) <= 500:
                raise ValueError('Describe tu participación (máximo 500 caracteres).')
            if len(opinion) > 2000:
                raise ValueError('Tu opinión debe tener como máximo 2000 caracteres.')
            if request.form.get('consent') != 'yes':
                raise ValueError('Autoriza el registro de tus datos para este evento.')
            if db.session.scalar(select(ExpoAttendance.id).where(func.lower(func.trim(ExpoAttendance.email)) == email)):
                flash('Este correo ya tiene asistencia registrada. No se sumó un registro nuevo.', 'warning')
                return redirect(url_for('expo.registration'))
            db.session.add(ExpoAttendance(name=name, email=email, category=category,
                                          participation=participation, opinion=opinion))
            try:
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                flash('Este correo ya tiene asistencia registrada. No se sumó un registro nuevo.', 'warning')
                return redirect(url_for('expo.registration'))
            # No account is created, and no submitted data is exposed in the URL.
            flash('¡Gracias! Tu asistencia y participación quedaron registradas.', 'success')
            return redirect(url_for('expo.registration', registrado='1'))
        except ValueError as error:
            flash(str(error), 'error')
    return render_template('expo_registration.html', categories=CATEGORIES, values=values,
                           registered=request.args.get('registrado') == '1')


def public_url():
    path = url_for('expo.registration')
    base = current_app.config.get('APP_BASE_URL', '')
    return base + path if base else url_for('expo.registration', _external=True)


def report_data():
    category = request.args.get('categoria', '')
    if category and category not in CATEGORIES:
        abort(400, description='Categoría inválida.')
    counts = dict(db.session.execute(select(ExpoAttendance.category, func.count()).group_by(ExpoAttendance.category)).all())
    query = select(ExpoAttendance).order_by(ExpoAttendance.recorded_at.desc(), ExpoAttendance.id.desc())
    if category:
        query = query.where(ExpoAttendance.category == category)
    return category, counts, query


@expo.get('/admin/expo-san-pablo')
@roles_required('admin')
def report():
    category, counts, query = report_data()
    page = db.paginate(query, per_page=25, max_per_page=25)
    return render_template('expo_report.html', categories=CATEGORIES, category=category,
                           counts=counts, total=sum(counts.values()), page=page, public_url=public_url())


@expo.get('/admin/expo-san-pablo/qr.png')
@roles_required('admin')
def qr():
    output = io.BytesIO()
    qrcode.make(public_url()).save(output, format='PNG')
    output.seek(0)
    return send_file(output, mimetype='image/png', as_attachment=request.args.get('descargar') == '1',
                     download_name='QR-Expo-San-Pablo.png')


def date_label(row):
    return local_datetime(row.recorded_at, current_app.config['APP_TIMEZONE']).strftime('%d/%m/%Y %H:%M')


@expo.get('/admin/expo-san-pablo/reporte.xlsx')
@roles_required('admin')
def excel_report():
    category, counts, query = report_data()
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Resumen'
    summary.append(['Expo San Pablo', 'Asistencia registrada'])
    summary.append(['Total general', sum(counts.values())])
    for key, label in CATEGORIES.items():
        summary.append([label, counts.get(key, 0)])
    summary.append(['Filtro del detalle', CATEGORIES.get(category, 'Todos')])
    summary.column_dimensions['A'].width = 30
    summary.column_dimensions['B'].width = 28
    sheet = workbook.create_sheet('Asistentes y opiniones')
    headers = ['Nombre', 'Correo', 'Categoría', 'Participación', 'Opinión', 'Fecha y hora (Guatemala)']
    sheet.append(headers)
    for row in db.session.scalars(query):
        sheet.append([row.name, row.email, CATEGORIES[row.category], row.participation, row.opinion, date_label(row)])
        # Treat attendee content as literal text, including leading '='.
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in zip('ABCDEF', [30, 38, 20, 55, 70, 28]):
        sheet.column_dimensions[index].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='Asistencia-Expo-San-Pablo.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@expo.get('/admin/expo-san-pablo/reporte.pdf')
@roles_required('admin')
def pdf_report():
    category, counts, query = report_data()
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    def paragraph(value, style='BodyText'):
        return Paragraph(escape(str(value)).replace('\n', '<br/>'), styles[style])
    story = [paragraph('Expo San Pablo — Reporte de asistencia', 'Title'),
             paragraph('Total general: ' + str(sum(counts.values())))]
    data = [['Categoría', 'Asistentes']] + [[label, str(counts.get(key, 0))] for key, label in CATEGORIES.items()]
    table = Table(data, colWidths=[320, 120])
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e9eef3')),
                               ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
    story += [Spacer(1, 12), table, Spacer(1, 18), paragraph('Detalle: ' + CATEGORIES.get(category, 'Todos'), 'Heading2')]
    found = False
    for row in db.session.scalars(query):
        found = True
        story += [paragraph(row.name, 'Heading3'), paragraph(row.email + ' · ' + CATEGORIES[row.category]),
                  paragraph('Registro: ' + date_label(row) + ' (Guatemala)'),
                  paragraph('Participación: ' + row.participation),
                  paragraph('Opinión: ' + (row.opinion or 'Sin opinión')), Spacer(1, 10)]
    if not found:
        story.append(paragraph('No hay asistentes registrados para esta selección.'))
    def footer(canvas, doc):
        canvas.setFont('Helvetica', 9)
        canvas.drawString(36, 22, 'Expo San Pablo · Página ' + str(doc.page))
    SimpleDocTemplate(output, pagesize=letter, leftMargin=36, rightMargin=36,
                      topMargin=36, bottomMargin=36).build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='Asistencia-Expo-San-Pablo.pdf', mimetype='application/pdf')
