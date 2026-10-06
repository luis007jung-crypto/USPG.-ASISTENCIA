import io
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func

from app import create_app, db
from app.models import ExpoAttendance, User


@pytest.fixture
def application():
    app = create_app({'TESTING': True, 'SECRET_KEY': 'test-only',
                      'SQLALCHEMY_DATABASE_URI': 'sqlite://', 'APP_BASE_URL': 'https://asistencia.example.org'})
    with app.app_context():
        db.create_all()
        admin = User(name='Admin', email='admin@administrador.uspg.edu.gt', role='admin')
        admin.set_password('test-password')
        student = User(name='Alumno', email='alumno@alumno.uspg.edu.gt', role='alumno')
        student.set_password('test-password')
        db.session.add_all([admin, student])
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


def submit(client, **changes):
    client.get('/expo-san-pablo')
    with client.session_transaction() as session:
        token = session['csrf_token']
    data = dict(csrf_token=token, name='Ana Pérez', email='ana@gmail.com', category='estudiante',
                participation='Visitante', opinion='Excelente evento', consent='yes')
    data.update(changes)
    return client.post('/expo-san-pablo', data=data, follow_redirects=True)


def sign_in(client, role):
    from flask import g
    g.pop('_login_user', None)
    user = db.session.scalar(select(User).where(User.role == role))
    with client.session_transaction() as session:
        session['_user_id'] = str(user.id)
        session['_fresh'] = True


@pytest.mark.parametrize('email,category', [('visitante@gmail.com','publico'), ('profe@outlook.com','catedratico'),
                                         ('estudiante@yahoo.com','estudiante'), ('alumno@alumno.uspg.edu.gt','estudiante')])
def test_public_registration_any_provider(application, email, category):
    client = application.test_client()
    assert client.get('/expo-san-pablo').status_code == 200
    response = submit(client, email=email, category=category)
    assert 'Registro completado' in response.text
    row = db.session.scalar(select(ExpoAttendance))
    assert (row.email, row.category, row.opinion) == (email, category, 'Excelente evento')
    assert db.session.scalar(select(func.count()).select_from(User)) == 2
    assert row.recorded_at is not None


def test_duplicate_and_csrf(application):
    client = application.test_client()
    submit(client)
    response = submit(client, email=' ANA@GMAIL.COM ', opinion='Reemplazo')
    assert 'ya tiene asistencia' in response.text
    assert db.session.scalar(select(func.count()).select_from(ExpoAttendance)) == 1
    assert db.session.scalar(select(ExpoAttendance.opinion)) == 'Excelente evento'
    assert client.post('/expo-san-pablo', data={'email': 'otra@gmail.com'}).status_code == 400


@pytest.mark.parametrize('changes', [{'email': 'correo inválido'}, {'category':'admin'}, {'consent':''},
                                   {'participation':''}, {'name':'x'}, {'opinion':'x'*2001}])
def test_invalid_input(application, changes):
    response = submit(application.test_client(), **changes)
    assert response.status_code == 200
    assert db.session.scalar(select(func.count()).select_from(ExpoAttendance)) == 0


@pytest.mark.parametrize('path', ['/admin/expo-san-pablo', '/admin/expo-san-pablo/qr.png',
                                 '/admin/expo-san-pablo/reporte.xlsx', '/admin/expo-san-pablo/reporte.pdf'])
def test_reports_protected(application, path):
    client = application.test_client()
    assert client.get(path).status_code == 302
    sign_in(client, 'alumno')
    assert client.get(path).status_code == 403


def test_report_counts_filters_exports_and_qr(application):
    client = application.test_client()
    for index, category in enumerate(['estudiante','catedratico','publico']):
        submit(client, email=f'asistente{index}@gmail.com', category=category,
               opinion='=SUM(A1:A2)' if index == 0 else '<script>alert(1)</script>')
    sign_in(client, 'admin')
    response = client.get('/admin/expo-san-pablo')
    assert response.status_code == 200
    assert '&lt;script&gt;' in response.text and '<script>alert(1)</script>' not in response.text
    from app.expo import report_data
    with application.test_request_context('/admin/expo-san-pablo'):
        category, counts, query = report_data()
        assert counts == {'estudiante':1,'catedratico':1,'publico':1}
    xlsx = client.get('/admin/expo-san-pablo/reporte.xlsx?categoria=estudiante')
    workbook = load_workbook(io.BytesIO(xlsx.data))
    assert workbook['Resumen']['B2'].value == 3
    detail = workbook['Asistentes y opiniones']
    assert detail.max_row == 2
    assert detail['E2'].value == '=SUM(A1:A2)' and detail['E2'].data_type == 's'
    pdf = client.get('/admin/expo-san-pablo/reporte.pdf')
    assert pdf.status_code == 200 and pdf.data.startswith(b'%PDF')
    with patch('app.expo.qrcode.make', wraps=__import__('qrcode').make) as make:
        first = client.get('/admin/expo-san-pablo/qr.png')
        second = client.get('/admin/expo-san-pablo/qr.png')
        assert first.data == second.data and first.data.startswith(b'\x89PNG')
        make.assert_called_with('https://asistencia.example.org/expo-san-pablo')
    assert client.get('/admin/expo-san-pablo?categoria=invalida').status_code == 400


def test_long_opinion_pdf_and_all_pages_exported(application):
    for index in range(30):
        db.session.add(ExpoAttendance(name=f'Asistente {index}', email=f'persona{index}@gmail.com',
                                      category='publico', participation='Visitante', opinion=('Muy bien. '*200)))
    db.session.commit()
    client = application.test_client()
    sign_in(client, 'admin')
    assert client.get('/admin/expo-san-pablo?page=2').status_code == 200
    workbook = load_workbook(io.BytesIO(client.get('/admin/expo-san-pablo/reporte.xlsx').data))
    assert workbook['Asistentes y opiniones'].max_row == 31
    assert client.get('/admin/expo-san-pablo/reporte.pdf').status_code == 200


def test_migration_idempotent(application):
    db.session.add(ExpoAttendance(name='Ana',email='ana@gmail.com',category='publico',participation='Visitante'))
    db.session.commit()
    runner = application.test_cli_runner()
    assert runner.invoke(args=['migrate-expo']).exit_code == 0
    assert runner.invoke(args=['migrate-expo']).exit_code == 0
    assert db.session.scalar(select(func.count()).select_from(ExpoAttendance)) == 1


def test_login_has_separate_public_expo_button(application):
    response = application.test_client().get('/')
    assert response.status_code == 200
    assert 'expo-entry-card' in response.text
    assert 'Registrar asistencia a la Expo' in response.text
    assert response.text.index('expo-entry-card') < response.text.index('name="identifier"')


def test_accounts_reject_duplicate_email_and_profile_change(application):
    from app.services import UserService
    service = UserService()
    service.register_student('Ana', 'ana@alumno.uspg.edu.gt', '2600002', 'Password123!')
    with pytest.raises(ValueError, match='cuenta con ese correo'):
        service.register_student('Otra Ana', ' ANA@ALUMNO.USPG.EDU.GT ', '2600003', 'Password123!')
    other = service.register_student('Luis', 'luis@alumno.uspg.edu.gt', '2600004', 'Password123!')
    with pytest.raises(ValueError):
        service.update_profile(other.id, 'Luis', ' ANA@ALUMNO.USPG.EDU.GT ', '2600004')
    assert db.session.scalar(select(func.count()).select_from(User)) == 4


def test_existing_account_can_attend_expo_only_once(application):
    client = application.test_client()
    email = 'alumno@alumno.uspg.edu.gt'
    submit(client, email=email, category='estudiante')
    response = submit(client, email=' ALUMNO@ALUMNO.USPG.EDU.GT ', category='publico')
    assert 'ya tiene asistencia' in response.text
    assert db.session.scalar(select(func.count()).select_from(ExpoAttendance)) == 1
    assert db.session.scalar(select(func.count()).select_from(User)) == 2


def test_legacy_mixed_case_duplicates_rejected(application):
    from app.services import UserService
    user = User(name='Ana', email='ANA@alumno.uspg.edu.gt', role='alumno', carnet='2600002')
    user.set_password('Password123!')
    db.session.add(user)
    db.session.add(ExpoAttendance(name='Ana', email='ANA@GMAIL.COM', category='publico', participation='Visitante'))
    db.session.commit()
    with pytest.raises(ValueError, match='cuenta con ese correo'):
        UserService().register_student('Ana', 'ana@alumno.uspg.edu.gt', '2600003', 'Password123!')
    assert 'ya tiene asistencia' in submit(application.test_client(), email='ana@gmail.com').text
